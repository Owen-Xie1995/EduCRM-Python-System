import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import mysql.connector
from openpyxl import load_workbook

from config import database_config
from database import DatabaseManager
from exports import export_students
from security import hash_password, hash_legacy_password, verify_password
from validation import validate_payment_amount, validate_billing_period


class ValidationTests(unittest.TestCase):
    def test_positive_integer_payments_only(self):
        for amount in (1, '1500', ' 80 '):
            self.assertGreater(validate_payment_amount(amount), 0)
        for amount in (0, -1, '0', '-2', '1.2', 1.2, '', None, True, '1e3', '100,000', 2**31):
            with self.subTest(amount=amount), self.assertRaises(ValueError):
                validate_payment_amount(amount)

    def test_month_and_year_must_be_valid(self):
        self.assertEqual(validate_billing_period('2026-10'), '2026-10')
        for period in ('2026-00', '2026-13', '26-10', '2026-1', '0000-01', '2026-10-01', None):
            with self.subTest(period=period), self.assertRaises(ValueError):
                validate_billing_period(period)


class PasswordTests(unittest.TestCase):
    def test_salted_roundtrip_preserves_spaces_and_unicode(self):
        password = ' 前後空白的密碼 '
        first, second = hash_password(password), hash_password(password)
        self.assertNotEqual(first, second)
        self.assertNotIn(password, first)
        self.assertTrue(verify_password(password, first))
        self.assertFalse(verify_password(password.strip(), first))
        self.assertFalse(verify_password('wrong', first))

    def test_plaintext_and_malformed_hashes_cannot_login(self):
        for value in ('1234', '', None, 'pbkdf2_sha256$1$aa$bb', 'pbkdf2_sha256$bad$xx$yy'):
            with self.subTest(value=value):
                self.assertFalse(verify_password('1234', value))

    def test_new_short_password_rejected_but_explicit_legacy_conversion_works(self):
        with self.assertRaises(ValueError):
            hash_password('1234')
        self.assertTrue(verify_password('1234', hash_legacy_password('1234')))


class ConfigTests(unittest.TestCase):
    @patch('config.load_dotenv')
    def test_requires_explicit_database_credentials(self, _):
        with patch.dict(os.environ, {}, clear=True), self.assertRaises(ValueError):
            database_config()

    @patch('config.load_dotenv')
    def test_uses_app_account_and_validates_port(self, _):
        with patch.dict(os.environ, {'DB_USER': 'educrm_app', 'DB_PASSWORD': 'local-test'}, clear=True):
            config = database_config()
            self.assertEqual(config['user'], 'educrm_app')
            self.assertEqual(config['port'], 3306)
            os.environ['DB_PORT'] = '65536'
            with self.assertRaises(ValueError):
                database_config()


class ExportTests(unittest.TestCase):
    def test_exports_visible_values_and_keeps_user_text_out_of_formulas(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'students.xlsx'
            export_students(path, [('S1', '=1+1', '女', '0000000000', '', '在學', '***', '2026-10-16')])
            workbook = load_workbook(path)
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet['B2'].value, '=1+1')
            self.assertEqual(sheet['B2'].data_type, 's')
            self.assertEqual(sheet['D2'].value, '0000000000')
            self.assertEqual(sheet['G2'].value, '***')
            workbook.close()


class StudentListTests(unittest.TestCase):
    def test_empty_search_does_not_reload_every_student(self):
        # 只呼叫畫面更新方法，無需建立 Tk 視窗或連線。
        from mydbsystemfinal import EduCRMApp
        screen = SimpleNamespace(tree=Mock(), db=Mock(), current_role='admin')
        screen.tree.get_children.return_value = ['old-row']
        EduCRMApp.load_student_data(screen, [])
        screen.db.fetch_all_students.assert_not_called()
        screen.tree.insert.assert_not_called()
        screen.tree.delete.assert_called_once_with('old-row')


class TeacherFormTests(unittest.TestCase):
    def form(self, callback, data=None, password=' teacher-password '):
        return SimpleNamespace(
            u_ent=Mock(get=Mock(return_value=' teacher02 ')),
            n_ent=Mock(get=Mock(return_value=' 陳老師 ')),
            p_ent=Mock(get=Mock(return_value=password)),
            callback=callback, data=data, destroy=Mock(),
        )

    @patch('mydbsystemfinal.messagebox.showerror')
    def test_short_new_password_shows_error_without_saving_or_closing(self, showerror):
        from mydbsystemfinal import TeacherFormDialog
        db = SimpleNamespace(cursor=Mock(), conn=Mock())
        form = self.form(lambda u, p, n: DatabaseManager.add_teacher(db, u, p, n), password='1234')
        TeacherFormDialog.save(form)
        showerror.assert_called_once_with('儲存失敗', '新密碼至少需要 8 個字元')
        db.cursor.execute.assert_not_called()
        form.destroy.assert_not_called()
        self.assertEqual(form.u_ent.get(), ' teacher02 ')
        self.assertEqual(form.p_ent.get(), '1234')

    @patch('mydbsystemfinal.messagebox.showerror')
    def test_duplicate_username_keeps_create_and_edit_forms_open(self, showerror):
        from mydbsystemfinal import TeacherFormDialog
        for data in (None, (7, 'old-account', '陳老師', 'staff')):
            with self.subTest(data=data):
                showerror.reset_mock()
                callback = Mock(side_effect=mysql.connector.IntegrityError('Duplicate entry', errno=1062))
                form = self.form(callback, data=data)
                TeacherFormDialog.save(form)
                showerror.assert_called_once_with('儲存失敗', '此帳號已被使用，請改用其他帳號。')
                form.destroy.assert_not_called()
                self.assertEqual(form.u_ent.get(), ' teacher02 ')

    @patch('mydbsystemfinal.messagebox.showerror')
    def test_success_strips_account_and_name_but_preserves_password(self, showerror):
        from mydbsystemfinal import TeacherFormDialog
        callback = Mock()
        form = self.form(callback)
        TeacherFormDialog.save(form)
        callback.assert_called_once_with('teacher02', ' teacher-password ', '陳老師')
        form.destroy.assert_called_once_with()
        showerror.assert_not_called()


if __name__ == '__main__':
    unittest.main()
